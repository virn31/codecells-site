# Genera docs/MODELO-FINANCIERO-*.xlsx (requiere openpyxl). Uso: python scripts/modelo-financiero.py salida.xlsx
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

OUT = sys.argv[1]
F = 'Arial'
BLUE = Font(name=F, color='0000FF')
BLACK = Font(name=F)
GREEN = Font(name=F, color='008000')
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
HEAD = Font(name=F, bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='1F3A5F')
SFILL = PatternFill('solid', fgColor='E8EEF5')
YFILL = PatternFill('solid', fgColor='FFFF00')
MXN = '$#,##0;($#,##0);-'
USD = '"US$"#,##0.00;("US$"#,##0.00);-'
NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'
PCT = '0.0%;(0.0%);-'
thin = Side(style='thin', color='BBBBBB')

wb = Workbook()
S = wb.active
S.title = 'Supuestos'
S.column_dimensions['A'].width = 58
for c in 'BCD':
    S.column_dimensions[c].width = 15
S.column_dimensions['E'].width = 70

S['A1'] = 'CODE CELLS™ — Modelo financiero 24 meses (dic-2026 a nov-2028)'
S['A1'].font = TITLE
S['A2'] = ('Cómo usarlo: cambia solo las celdas en AZUL (las amarillas son las decisiones clave). '
           'Todo lo demás se calcula. Montos en MXN salvo que diga USD. Generado 2026-10-19.')
S['A2'].font = Font(name=F, italic=True)
S['A2'].alignment = Alignment(wrap_text=True)
S.merge_cells('A2:E2')
S.row_dimensions[2].height = 30

R = {}  # nombre -> fila


def seccion(fila, texto, escenarios=False):
    S.cell(fila, 1, texto).font = HEAD
    for c in range(1, 6):
        S.cell(fila, c).fill = HFILL
    if escenarios:
        for i, t in enumerate(['Conservador', 'Base', 'Optimista']):
            S.cell(fila, 2 + i, t).font = HEAD
            S.cell(fila, 2 + i).alignment = Alignment(horizontal='center')
    S.cell(fila, 5, 'Fuente / nota').font = HEAD


def entrada(fila, clave, etiqueta, valores, fmt, nota, clave_dec=False):
    R[clave] = fila
    S.cell(fila, 1, etiqueta).font = BLACK
    if not isinstance(valores, list):
        valores = [valores]
    for i, v in enumerate(valores):
        c = S.cell(fila, 2 + i, v)
        c.font = BLUE
        c.number_format = fmt
        if clave_dec:
            c.fill = YFILL
    S.cell(fila, 5, nota).font = Font(name=F, size=9)
    S.cell(fila, 5).alignment = Alignment(wrap_text=True, vertical='top')


f = 4
seccion(f, 'Precios (MXN, con IVA)'); f += 1
entrada(f, 'p_red', 'Precio plan Red (mensual)', 600, MXN, 'Dato de Víctor 2026-10-19. Directorio + solicitudes de cita.', True); f += 1
entrada(f, 'p_clin', 'Precio plan Clínico (mensual, incluye todo)', 1100, MXN, 'Dato de Víctor 2026-10-19. Red + portal clínico + NOVA + app del paciente.', True); f += 1
entrada(f, 'p_fund', 'Precio fundador mientras solo existe Red (mensual)', 600, MXN, 'Propuesta: el fundador paga lo mismo que Red mientras el portal clínico no existe.', True); f += 1
entrada(f, 'p_fund_clin', 'Precio fundador con portal clínico (mensual)', 800, MXN, 'Propuesta: al lanzarse el Clínico el fundador lo recibe a este precio congelado 24 meses (36 los primeros 25), contra 1 100 de lista. A 600 su contribución cae a ~147 MXN (28%). En 24 meses ningún fundador sale del precio congelado.', True); f += 1
entrada(f, 'lug_fund', 'Lugares de fundador', 100, NUM, 'Oferta de 100 médicos fundadores.'); f += 1
entrada(f, 'iva', 'IVA', 0.16, PCT, 'Los ingresos del modelo se calculan SIN IVA (el IVA se entera al SAT).'); f += 1
entrada(f, 'p_mig', 'Precio de migración por paciente (con IVA)', 58, MXN, 'Propuesta (Víctor decidió que la migración se cobra aparte): ~50 + IVA por paciente. Costo estimado ~32 por paciente (ver bloque Migración).', True); f += 1

f += 1
seccion(f, 'Tipo de cambio y costos en USD'); f += 1
entrada(f, 'tc', 'Tipo de cambio MXN por USD', 18.49, '0.00', 'FIX Banxico 16-oct-2026 (18.4930). Prueba con 21 para ver el efecto de una devaluación.'); f += 1
entrada(f, 'infra0', 'Infraestructura actual (USD/mes)', 68.33, USD, 'Dato de Víctor: Airtable 20 + Vercel 20 + API Anthropic 20 + dominio 100/año.'); f += 1
entrada(f, 'infra1', 'Infraestructura nueva desde el lanzamiento Clínico (USD/mes)', 80, USD, 'Estimado: PostgreSQL ~25, respaldos, monitoreo 0–26, correo 0–20, Cloudflare R2. Ver docs/VALUACION-2026-10-19.md §4.'); f += 1
entrada(f, 'ia_red', 'IA por médico en plan Red (USD/mes)', 0.5, USD, 'Estimado: el directorio casi no usa NOVA.'); f += 1
entrada(f, 'ia_clin', 'IA por médico en plan Clínico (USD/mes)', 14.1, USD, 'Estimado con caché de instrucciones: mezcla 30% bajo (4) / 50% medio (13) / 20% alto (32). SIN caché se duplica (~28). Tope de NOVA por paciente (decisión de Víctor).', True); f += 1

f += 1
seccion(f, 'Almacenamiento de archivos (por GB)'); f += 1
entrada(f, 'gb_inc', 'GB incluidos en el plan Clínico', 20, NUM, 'Propuesta. Ejemplo de Víctor: 4–6 fotos pre/post (~4 MB c/u) + consentimientos ≈ 26 MB por paciente quirúrgico → 100 pacientes al año ≈ 2.6 GB.', True); f += 1
entrada(f, 'gb_inc_red', 'GB incluidos en el plan Red', 1, NUM, 'Propuesta: solo fotos de perfil y documentos del directorio.'); f += 1
entrada(f, 'p_gb', 'Precio por GB adicional al mes (con IVA)', 2, MXN, 'Propuesta. Para volúmenes de TB (imagenología) ofrecer paquete, p. ej. 1 TB por ~$990/mes (costo ~$410).', True); f += 1
entrada(f, 'gb_mes', 'GB nuevos por médico Clínico al mes (promedio)', 0.5, '0.00', 'Estimado: mezcla de especialidades. Muy variable: de MB a TB según la especialidad.', True); f += 1
entrada(f, 'pct_heavy', '% de médicos que superan lo incluido', 0.10, PCT, 'Estimado: cirugía plástica, dermatología, imagen, oftalmología.', True); f += 1
entrada(f, 'gb_heavy', 'GB adicionales promedio de esos médicos', 100, NUM, 'Estimado. Un médico de imagenología puede llegar a TB: cotizar paquete.', True); f += 1
entrada(f, 'c_gb', 'Costo por GB al mes, copia principal (USD)', 0.015, '"US$"0.000', 'Verificado: Cloudflare R2 0.015 USD/GB, sin cargo por descarga.'); f += 1
entrada(f, 'c_gb_resp', 'Costo por GB al mes, copia de respaldo (USD)', 0.00695, '"US$"0.00000', 'Secundario: Backblaze B2 ~6.95 USD/TB. Respaldo en otro proveedor.'); f += 1

f += 1
seccion(f, 'Cobro'); f += 1
entrada(f, 'com_pct', 'Comisión de la pasarela (% del cobro)', 0.036, PCT, 'Secundario: Stripe México ~3.6% + 3 MXN por tarjeta nacional. Se suma IVA sobre la comisión.'); f += 1
entrada(f, 'com_fija', 'Comisión fija por cobro (MXN)', 3, MXN, 'Secundario (Stripe).'); f += 1
entrada(f, 'cfdi', 'Timbrado CFDI por factura (MXN)', 2, MXN, 'Secundario: 0.50–2 MXN por folio según el PAC.'); f += 1

f += 1
seccion(f, 'Operación'); f += 1
entrada(f, 'contador', 'Contador (MXN/mes)', 2500, MXN, 'COTIZAR. Valor de ejemplo.', True); f += 1
entrada(f, 'sop_h', 'Horas de soporte por médico al mes', 0.5, NUM1, 'Estimado.'); f += 1
entrada(f, 'sop_c', 'Costo por hora de soporte (MXN)', 150, MXN, 'Estimado (asistente de soporte).'); f += 1
entrada(f, 'cac', 'Costo de adquisición por médico nuevo (MXN)', 1500, MXN, 'Estimado: publicidad, visitas, demo. Medirlo en el piloto.', True); f += 1
entrada(f, 'sueldo', 'Sueldo de Víctor (MXN/mes)', 0, MXN, 'Decisión: en 0 el modelo muestra el flujo sin pagar tu tiempo. Ponle un monto para ver cuándo se sostiene pagándote.', True); f += 1
entrada(f, 'abogado', 'Abogado (pago único, mes 1)', 40000, MXN, 'COTIZAR: aviso de privacidad clínico, contrato de afiliación, términos.', True); f += 1
entrada(f, 'pentest', 'Prueba de penetración (pago único, mes de lanzamiento Clínico)', 50000, MXN, 'COTIZAR.', True); f += 1
entrada(f, 'mes_clin', 'Mes de lanzamiento del plan Clínico (1 = dic-2026)', 6, NUM, 'Supuesto: mayo-2027, tras acta constitutiva + IMPI + abogado.', True); f += 1

f += 1
seccion(f, 'Migración de expedientes (costo de producirla)'); f += 1
entrada(f, 'mig_pac', 'Pacientes promedio por migración', 200, NUM, 'Estimado.'); f += 1
entrada(f, 'mig_pag', 'Páginas por paciente', 10, NUM, 'Estimado.'); f += 1
entrada(f, 'mig_ia', 'IA por página (USD)', 0.013, '"US$"0.000', 'Estimado (lectura de imagen con Claude Sonnet 5; la mitad en modo por lotes).'); f += 1
entrada(f, 'mig_min', 'Minutos de revisión humana por página', 1.5, NUM1, 'Estimado: verificar identidad y valores (no se inventa nada).'); f += 1
entrada(f, 'mig_hora', 'Costo por hora de revisión (MXN)', 120, MXN, 'Estimado.'); f += 1

f += 1
seccion(f, 'Canal Regene Global'); f += 1
entrada(f, 'rg_total', 'Médicos afiliados a Regene Global', 2500, NUM, 'Dato de Víctor 2026-10-19 (aproximado).'); f += 1
entrada(f, 'rg_inicio', 'Mes en que arranca el canal Regene (1 = dic-2026)', 7, NUM, 'Supuesto: un mes después del lanzamiento Clínico, ya con la base clínica fuera de Airtable.', True); f += 1
entrada(f, 'rg_desc', 'Descuento o comisión para el canal Regene', 0.15, PCT, 'Propuesta: 15% del precio Clínico (descuento al médico o comisión a Regene). Decisión de negocio.', True); f += 1
entrada(f, 'rg_cac', 'Costo de incorporar a un médico Regene (MXN)', 300, MXN, 'Estimado: alta, capacitación grupal. Menor que el costo de adquisición normal.'); f += 1
entrada(f, 'infra_100', 'Infraestructura adicional por cada 100 médicos (USD/mes)', 25, USD, 'Estimado: base de datos y cómputo más grandes al crecer.'); f += 1

f += 1
seccion(f, 'Escenarios', escenarios=True); f += 1
entrada(f, 'nuevos', 'Médicos nuevos por mes', [3, 6, 10], NUM, 'Estimado. El piloto de diciembre es de 5–10 médicos.', True); f += 1
entrada(f, 'churn', 'Cancelación mensual', [0.04, 0.03, 0.02], PCT, 'Estimado. SaaS para médicos pequeños: 2–5% mensual es un rango común.', True); f += 1
entrada(f, 'pct_clin', '% de médicos nuevos (no fundadores) que eligen Clínico tras el lanzamiento', [0.4, 0.5, 0.6], PCT, 'Estimado.'); f += 1
entrada(f, 'upg', '% de médicos Red que suben a Clínico cada mes tras el lanzamiento', [0.03, 0.05, 0.08], PCT, 'Estimado.'); f += 1
entrada(f, 'pct_mig', '% de médicos que entran a Clínico y compran migración', [0.2, 0.3, 0.4], PCT, 'Estimado.'); f += 1
entrada(f, 'rg_pct', '% de médicos Regene que se afilian (repartido del mes de arranque al mes 24)', [0.05, 0.10, 0.20], PCT, 'Estimado. 5% = 125 médicos; 20% = 500.', True); f += 1
ULT_SUP = f

MESES = ['dic-2026'] + [f'{m}-2027' for m in ['ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep', 'oct', 'nov', 'dic']] + \
        [f'{m}-2028' for m in ['ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep', 'oct', 'nov']]
assert len(MESES) == 24


def s(clave, col='B'):
    return f'Supuestos!${col}${R[clave]}'


FILAS = {}


def hoja(nombre, colesc):
    W = wb.create_sheet(nombre)
    W.column_dimensions['A'].width = 46
    for i in range(24):
        W.column_dimensions[L(3 + i)].width = 11
    W.column_dimensions['B'].width = 3
    W['A1'] = f'Escenario {nombre} — mes a mes (MXN sin IVA)'
    W['A1'].font = TITLE
    W.freeze_panes = 'C5'
    filas = [
        ('mes', 'Mes #'), ('etq', 'Mes'), ('lanz', 'Plan Clínico disponible (1 = sí)'), None,
        ('H', 'MÉDICOS'),
        ('nuevos', 'Médicos nuevos'), ('fn', 'Fundadores nuevos'), ('fcum', 'Fundadores inscritos (acumulado)'),
        ('nf', 'Nuevos no fundadores'), ('nclin', 'Nuevos que eligen Clínico'), ('nred', 'Nuevos que eligen Red'),
        ('fund', 'Fundadores activos'), ('upg', 'Suben de Red a Clínico'), ('red', 'Red activos'),
        ('clin', 'Clínico activos (precio completo)'), ('rgn', 'Médicos Regene nuevos'), ('rg', 'Médicos Regene activos (Clínico con descuento)'),
        ('tot', 'Total médicos activos'),
        ('usoclin', 'Médicos usando el portal clínico'), ('usored', 'Médicos usando solo Red'), None,
        ('H', 'INGRESOS'),
        ('i_fund', 'Fundadores'), ('i_red', 'Plan Red'), ('i_clin', 'Plan Clínico'), ('i_rg', 'Canal Regene'),
        ('gb_ext', 'GB adicionales cobrados (médicos de alto consumo)'), ('gb', 'GB guardados en total (se conservan 5 años, NOM-004)'), ('i_alm', 'Almacenamiento adicional'),
        ('mrr', 'Ingreso recurrente mensual (MRR)'),
        ('migs', 'Migraciones vendidas'), ('i_mig', 'Ingreso por migraciones'), ('ing', 'INGRESO TOTAL'), None,
        ('H', 'COSTOS VARIABLES'),
        ('c_ia', 'IA (NOVA)'), ('c_alm', 'Almacenamiento (2 copias)'), ('c_com', 'Comisión de cobro (con IVA de la comisión)'), ('c_cfdi', 'Facturas CFDI'),
        ('c_sop', 'Soporte'), ('c_mig', 'Producción de migraciones'), ('c_var', 'Total costos variables'),
        ('mc', 'MARGEN DE CONTRIBUCIÓN'), ('mcp', 'Margen de contribución %'), None,
        ('H', 'COSTOS FIJOS Y DE CRECIMIENTO'),
        ('c_inf', 'Infraestructura'), ('c_cont', 'Contador'), ('c_cac', 'Adquisición de médicos'), ('c_sue', 'Sueldo de Víctor'),
        ('c_uni', 'Pagos únicos (abogado, prueba de penetración)'), ('c_fij', 'Total fijos y crecimiento'), None,
        ('flujo', 'FLUJO NETO DEL MES'), ('acum', 'FLUJO ACUMULADO'), ('arr', 'Ingreso recurrente anualizado (ARR)'),
        ('pos', 'Ayuda: mes con flujo positivo (99 = no)'),
    ]
    r = 3
    fr = {}
    for item in filas:
        if item is None:
            r += 1
            continue
        k, t = item
        if k == 'H':
            W.cell(r, 1, t).font = BOLD
            for c in range(1, 27):
                W.cell(r, c).fill = SFILL
        else:
            fr[k] = r
            W.cell(r, 1, t).font = BOLD if t.isupper() or k in ('mrr', 'flujo', 'acum') else BLACK
        r += 1
    FILAS[nombre] = fr
    e = lambda clave: s(clave, colesc)
    for i in range(24):
        c = L(3 + i)
        p = L(2 + i)  # columna anterior
        primero = i == 0
        V = {}
        V['mes'] = i + 1
        V['etq'] = MESES[i]
        V['lanz'] = f'=IF({c}{fr["mes"]}>={s("mes_clin")},1,0)'
        V['nuevos'] = f'={e("nuevos")}'
        V['fn'] = f'=MIN({c}{fr["nuevos"]},{s("lug_fund")})' if primero else f'=MIN({c}{fr["nuevos"]},{s("lug_fund")}-{p}{fr["fcum"]})'
        V['fcum'] = f'={c}{fr["fn"]}' if primero else f'={p}{fr["fcum"]}+{c}{fr["fn"]}'
        V['nf'] = f'={c}{fr["nuevos"]}-{c}{fr["fn"]}'
        V['nclin'] = f'={c}{fr["nf"]}*{c}{fr["lanz"]}*{e("pct_clin")}'
        V['nred'] = f'={c}{fr["nf"]}-{c}{fr["nclin"]}'
        ret = f'(1-{e("churn")})'
        V['fund'] = f'={c}{fr["fn"]}' if primero else f'={p}{fr["fund"]}*{ret}+{c}{fr["fn"]}'
        V['upg'] = '=0' if primero else f'={p}{fr["red"]}*{ret}*{c}{fr["lanz"]}*{e("upg")}'
        V['red'] = f'={c}{fr["nred"]}' if primero else f'={p}{fr["red"]}*{ret}-{c}{fr["upg"]}+{c}{fr["nred"]}'
        V['clin'] = f'={c}{fr["nclin"]}' if primero else f'={p}{fr["clin"]}*{ret}+{c}{fr["nclin"]}+{c}{fr["upg"]}'
        V['rgn'] = f'=IF({c}{fr["mes"]}>={s("rg_inicio")},{s("rg_total")}*{e("rg_pct")}/MAX(1,25-{s("rg_inicio")}),0)'
        V['rg'] = f'={c}{fr["rgn"]}' if primero else f'={p}{fr["rg"]}*{ret}+{c}{fr["rgn"]}'
        V['tot'] = f'={c}{fr["fund"]}+{c}{fr["red"]}+{c}{fr["clin"]}+{c}{fr["rg"]}'
        V['usoclin'] = f'={c}{fr["clin"]}+{c}{fr["rg"]}+{c}{fr["fund"]}*{c}{fr["lanz"]}'
        V['usored'] = f'={c}{fr["tot"]}-{c}{fr["usoclin"]}'
        iv = f'(1+{s("iva")})'
        V['i_fund'] = f'={c}{fr["fund"]}*IF({c}{fr["lanz"]}=1,{s("p_fund_clin")},{s("p_fund")})/{iv}'
        V['i_red'] = f'={c}{fr["red"]}*{s("p_red")}/{iv}'
        V['i_clin'] = f'={c}{fr["clin"]}*{s("p_clin")}/{iv}'
        V['i_rg'] = f'={c}{fr["rg"]}*{s("p_clin")}*(1-{s("rg_desc")})/{iv}'
        V['gb_ext'] = f'={c}{fr["usoclin"]}*{s("pct_heavy")}*{s("gb_heavy")}'
        V['gb'] = (f'={c}{fr["usoclin"]}*{s("gb_mes")}' if primero else f'={p}{fr["gb"]}-{p}{fr["gb_ext"]}+{c}{fr["usoclin"]}*{s("gb_mes")}') + f'+{c}{fr["gb_ext"]}'
        V['i_alm'] = f'={c}{fr["gb_ext"]}*{s("p_gb")}/{iv}'
        V['mrr'] = f'={c}{fr["i_fund"]}+{c}{fr["i_red"]}+{c}{fr["i_clin"]}+{c}{fr["i_rg"]}+{c}{fr["i_alm"]}'
        V['migs'] = f'=({c}{fr["nclin"]}+{c}{fr["upg"]}+{c}{fr["rgn"]})*{e("pct_mig")}'
        V['i_mig'] = f'={c}{fr["migs"]}*{s("mig_pac")}*{s("p_mig")}/{iv}'
        V['ing'] = f'={c}{fr["mrr"]}+{c}{fr["i_mig"]}'
        V['c_ia'] = f'=({c}{fr["usored"]}*{s("ia_red")}+{c}{fr["usoclin"]}*{s("ia_clin")})*{s("tc")}'
        V['c_alm'] = f'={c}{fr["gb"]}*({s("c_gb")}+{s("c_gb_resp")})*{s("tc")}'
        V['c_com'] = f'=(({c}{fr["tot"]}+{c}{fr["migs"]})*{s("com_fija")}+{c}{fr["ing"]}*{iv}*{s("com_pct")})*{iv}'
        V['c_cfdi'] = f'=({c}{fr["tot"]}+{c}{fr["migs"]})*{s("cfdi")}'
        V['c_sop'] = f'={c}{fr["tot"]}*{s("sop_h")}*{s("sop_c")}'
        V['c_mig'] = f'={c}{fr["migs"]}*{s("mig_pac")}*{s("mig_pag")}*({s("mig_ia")}*{s("tc")}+{s("mig_min")}/60*{s("mig_hora")})'
        V['c_var'] = f'=SUM({c}{fr["c_ia"]}:{c}{fr["c_mig"]})'
        V['mc'] = f'={c}{fr["ing"]}-{c}{fr["c_var"]}'
        V['mcp'] = f'=IF({c}{fr["ing"]}=0,0,{c}{fr["mc"]}/{c}{fr["ing"]})'
        V['c_inf'] = f'=({s("infra0")}+{s("infra1")}*{c}{fr["lanz"]}+{s("infra_100")}*INT({c}{fr["tot"]}/100))*{s("tc")}'
        V['c_cont'] = f'={s("contador")}'
        V['c_cac'] = f'={c}{fr["nuevos"]}*{s("cac")}+{c}{fr["rgn"]}*{s("rg_cac")}'
        V['c_sue'] = f'={s("sueldo")}'
        V['c_uni'] = f'=IF({c}{fr["mes"]}=1,{s("abogado")},0)+IF({c}{fr["mes"]}={s("mes_clin")},{s("pentest")},0)'
        V['c_fij'] = f'=SUM({c}{fr["c_inf"]}:{c}{fr["c_uni"]})'
        V['flujo'] = f'={c}{fr["mc"]}-{c}{fr["c_fij"]}'
        V['acum'] = f'={c}{fr["flujo"]}' if primero else f'={p}{fr["acum"]}+{c}{fr["flujo"]}'
        V['arr'] = f'={c}{fr["mrr"]}*12'
        V['pos'] = f'=IF({c}{fr["flujo"]}>0,{c}{fr["mes"]},99)'
        for k, v in V.items():
            cell = W.cell(fr[k], 3 + i, v)
            cell.font = BLUE if k == 'mes' else (GREEN if isinstance(v, str) and v.startswith('=') and 'Supuestos!' in v and v.count('!') == 1 and k in ('nuevos', 'c_cont', 'c_sue') else BLACK)
            if k in ('mes', 'etq'):
                cell.font = BOLD
                cell.alignment = Alignment(horizontal='center')
            elif k in ('lanz', 'pos'):
                cell.number_format = '0'
            elif k in ('nuevos', 'fn', 'fcum', 'nf', 'nclin', 'nred', 'fund', 'upg', 'red', 'clin', 'rgn', 'rg', 'tot', 'usoclin', 'usored', 'migs', 'gb', 'gb_ext'):
                cell.number_format = NUM1
            elif k == 'mcp':
                cell.number_format = PCT
            else:
                cell.number_format = MXN
            if k in ('mrr', 'ing', 'mc', 'flujo', 'acum'):
                cell.font = Font(name=F, bold=True)
    for k in ('mrr', 'ing', 'mc', 'flujo', 'acum'):
        for cc in range(1, 27):
            W.cell(fr[k], cc).border = Border(top=thin)
    W.cell(fr['mes'], 1).comment = Comment('Los médicos se muestran con decimales porque la cancelación es un promedio; léelos como "aprox."', 'Claude')
    return W


ESC = [('Conservador', 'B'), ('Base', 'C'), ('Optimista', 'D')]
for n, col in ESC:
    hoja(n, col)

# ── Resumen ──
X = wb.create_sheet('Resumen', 0)
X.column_dimensions['A'].width = 60
for c in 'BCD':
    X.column_dimensions[c].width = 16
X.column_dimensions['E'].width = 60
X['A1'] = 'CODE CELLS™ — Resumen del modelo a 12 y 24 meses'
X['A1'].font = TITLE
X['A2'] = 'MXN sin IVA. Todo se calcula desde la hoja Supuestos. Estimaciones: medir y corregir con el piloto.'
X['A2'].font = Font(name=F, italic=True)
r = 4
X.cell(r, 1, 'Resultado').font = HEAD
for i, (n, _) in enumerate(ESC):
    X.cell(r, 2 + i, n).font = HEAD
for c in range(1, 6):
    X.cell(r, c).fill = HFILL
X.cell(r, 5, 'Cómo leerlo').font = HEAD
r += 1
COL12, COL24 = L(3 + 11), L(3 + 23)
fila_fin = lambda n, k: FILAS[n][k]
items = [
    ('Médicos activos al mes 12 (nov-2027)', lambda n: f"='{n}'!{COL12}{fila_fin(n,'tot')}", NUM, ''),
    ('Médicos activos al mes 24 (nov-2028)', lambda n: f"='{n}'!{COL24}{fila_fin(n,'tot')}", NUM, ''),
    ('   de ellos fundadores', lambda n: f"='{n}'!{COL24}{fila_fin(n,'fund')}", NUM, 'Los fundadores pagan el precio congelado.'),
    ('   de ellos por el canal Regene', lambda n: f"='{n}'!{COL24}{fila_fin(n,'rg')}", NUM, 'Clínico con descuento/comisión del canal.'),
    ('GB guardados al mes 24', lambda n: f"='{n}'!{COL24}{fila_fin(n,'gb')}", NUM, 'Se acumulan: la NOM-004 obliga a conservar 5 años aunque el médico se vaya.'),
    ('Ingreso por almacenamiento adicional, mes 24', lambda n: f"='{n}'!{COL24}{fila_fin(n,'i_alm')}", MXN, ''),
    ('Costo de almacenamiento (2 copias), mes 24', lambda n: f"='{n}'!{COL24}{fila_fin(n,'c_alm')}", MXN, ''),
    ('Ingreso recurrente mensual (MRR) mes 12', lambda n: f"='{n}'!{COL12}{fila_fin(n,'mrr')}", MXN, ''),
    ('Ingreso recurrente mensual (MRR) mes 24', lambda n: f"='{n}'!{COL24}{fila_fin(n,'mrr')}", MXN, ''),
    ('Ingreso recurrente anualizado (ARR) mes 24', lambda n: f"='{n}'!{COL24}{fila_fin(n,'arr')}", MXN, 'MRR × 12.'),
    ('Ingreso total acumulado 24 meses', lambda n: f"=SUM('{n}'!C{fila_fin(n,'ing')}:{COL24}{fila_fin(n,'ing')})", MXN, ''),
    ('Margen de contribución mes 24', lambda n: f"='{n}'!{COL24}{fila_fin(n,'mcp')}", PCT, 'Lo que queda de cada peso tras IA, cobro, soporte y migraciones.'),
    ('Flujo acumulado al mes 12', lambda n: f"='{n}'!{COL12}{fila_fin(n,'acum')}", MXN, ''),
    ('Flujo acumulado al mes 24', lambda n: f"='{n}'!{COL24}{fila_fin(n,'acum')}", MXN, 'Negativo = todavía no recupera lo invertido.'),
    ('Capital requerido (peor punto del flujo acumulado)', lambda n: f"=MAX(0,-MIN('{n}'!C{fila_fin(n,'acum')}:{COL24}{fila_fin(n,'acum')}))", MXN, 'El dinero que hay que tener para no quedarse sin caja.'),
    ('Primer mes con flujo positivo (# de mes)', lambda n: f"=MIN('{n}'!C{fila_fin(n,'pos')}:{COL24}{fila_fin(n,'pos')})", '0', '99 = no ocurre en 24 meses. Mes 1 = dic-2026.'),
]
for t, fn, fmt, nota in items:
    X.cell(r, 1, t).font = BLACK
    for i, (n, _) in enumerate(ESC):
        c = X.cell(r, 2 + i, fn(n))
        c.font = GREEN
        c.number_format = fmt
    X.cell(r, 5, nota).font = Font(name=F, size=9)
    r += 1

r += 1
X.cell(r, 1, 'Economía por médico al mes (estado estable)').font = HEAD
for c in range(1, 6):
    X.cell(r, c).fill = HFILL
for i, t in enumerate(['Red', 'Clínico', 'Fundador*']):
    X.cell(r, 2 + i, t).font = HEAD
r += 1
iv = f"(1+{s('iva')})"
ue_ing = r
X.cell(r, 1, 'Ingreso sin IVA')
X.cell(r, 2, f"={s('p_red')}/{iv}"); X.cell(r, 3, f"={s('p_clin')}/{iv}"); X.cell(r, 4, f"={s('p_fund_clin')}/{iv}")
r += 1
ue_ia = r
X.cell(r, 1, 'IA (NOVA)')
X.cell(r, 2, f"={s('ia_red')}*{s('tc')}"); X.cell(r, 3, f"={s('ia_clin')}*{s('tc')}"); X.cell(r, 4, f"=C{r}")
r += 1
ue_com = r
X.cell(r, 1, 'Comisión de cobro + factura')
for i, k in enumerate(['p_red', 'p_clin', 'p_fund_clin']):
    X.cell(r, 2 + i, f"=({s('com_fija')}+{s(k)}*{s('com_pct')})*{iv}+{s('cfdi')}")
r += 1
ue_sop = r
X.cell(r, 1, 'Soporte')
for i in range(3):
    X.cell(r, 2 + i, f"={s('sop_h')}*{s('sop_c')}")
r += 1
ue_mc = r
X.cell(r, 1, 'Contribución por médico').font = BOLD
for i in range(3):
    c = L(2 + i)
    X.cell(r, 2 + i, f"={c}{ue_ing}-{c}{ue_ia}-{c}{ue_com}-{c}{ue_sop}").font = BOLD
r += 1
X.cell(r, 1, 'Contribución %')
for i in range(3):
    c = L(2 + i)
    X.cell(r, 2 + i, f"=IF({c}{ue_ing}=0,0,{c}{ue_mc}/{c}{ue_ing})").number_format = PCT
r += 1
for rr in range(ue_ing, ue_mc + 1):
    for i in range(3):
        X.cell(rr, 2 + i).number_format = MXN
        if X.cell(rr, 2 + i).font != BOLD:
            X.cell(rr, 2 + i).font = BLACK
X.cell(r, 1, '* Fundador ya con portal clínico (después del lanzamiento).').font = Font(name=F, size=9, italic=True)
r += 2
X.cell(r, 1, 'Punto de equilibrio').font = HEAD
for c in range(1, 6):
    X.cell(r, c).fill = HFILL
r += 1
fij = r
X.cell(r, 1, 'Costos fijos mensuales tras el lanzamiento (infra + contador + sueldo)')
X.cell(r, 2, f"=({s('infra0')}+{s('infra1')})*{s('tc')}+{s('contador')}+{s('sueldo')}").number_format = MXN
r += 1
X.cell(r, 1, 'Médicos Clínico necesarios para cubrir los fijos')
X.cell(r, 2, f"=IF(C{ue_mc}<=0,\"No se cubre\",ROUNDUP(B{fij}/C{ue_mc},0))").number_format = '0'
X.cell(r, 5, 'Sin contar adquisición ni pagos únicos.').font = Font(name=F, size=9)
r += 1
X.cell(r, 1, 'Médicos fundadores necesarios para cubrir los fijos')
X.cell(r, 2, f"=IF(D{ue_mc}<=0,\"No se cubre\",ROUNDUP(B{fij}/D{ue_mc},0))").number_format = '0'
r += 1
X.cell(r, 1, 'Médicos Clínico necesarios para pagarte un sueldo de')
X.cell(r, 2, 40000).font = BLUE
X.cell(r, 2).number_format = MXN
X.cell(r, 2).fill = YFILL
X.cell(r, 3, f"=IF(C{ue_mc}<=0,\"No se cubre\",ROUNDUP((B{fij}-{s('sueldo')}+B{r})/C{ue_mc},0))").number_format = '0'
X.cell(r, 5, 'Cambia el monto en azul. Resultado en la columna C.').font = Font(name=F, size=9)

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if cell.font is None or cell.font.name != F:
                cell.font = Font(name=F, bold=cell.font.bold if cell.font else False,
                                 color=cell.font.color if cell.font else None,
                                 italic=cell.font.italic if cell.font else False,
                                 size=cell.font.size if cell.font else 11)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('ok', OUT)
